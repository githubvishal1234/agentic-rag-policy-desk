from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.models.document import Document


def load_txt(path: Path) -> list[Document]:
    """Load a TXT document."""

    text = path.read_text(encoding="utf-8")

    return [
        Document(
            text=text,
            metadata={
                "source": path.name,
                "file_type": "txt",
            },
        )
    ]


def load_pdf(path: Path) -> list[Document]:
    """Load a PDF document page by page."""

    reader = PdfReader(str(path))

    documents = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        if text.strip():
            documents.append(
                Document(
                    text=text,
                    metadata={
                        "source": path.name,
                        "file_type": "pdf",
                        "page": page_number,
                    },
                )
            )

    return documents


def load_docx(path: Path) -> list[Document]:
    """Load a DOCX document."""

    docx = DocxDocument(str(path))

    paragraphs = []

    for paragraph in docx.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    text = "\n".join(paragraphs)

    return [
        Document(
            text=text,
            metadata={
                "source": path.name,
                "file_type": "docx",
            },
        )
    ]


def load_document(file_path: str) -> list[Document]:
    """Load a supported document."""

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    extension = path.suffix.lower()

    if extension == ".txt":
        return load_txt(path)

    if extension == ".pdf":
        return load_pdf(path)

    if extension == ".docx":
        return load_docx(path)

    raise ValueError(
        f"Unsupported file type: {extension}. "
        "Supported types: .txt, .pdf, .docx"
    )